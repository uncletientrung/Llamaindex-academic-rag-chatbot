
import asyncio
import json
from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from llama_index.core.evaluation import RetrieverEvaluator

from test import RAGService


DATASET_PATH = Path(
    "evaluation/retrieval_evaluation/dataset.json"
)

RESULT_PATH = Path(
    "evaluation/retrieval_evaluation/results/retrieval.xlsx"
)

METRICS = ["hit_rate", "mrr", "precision"]


def build_evaluator(rag: RAGService):
    """Tạo evaluator sử dụng retriever và postprocessor hiện tại."""

    if rag.router_retriever is None:
        raise ValueError(
            "RouterRetriever chưa được khởi tạo. "
            "Hãy kiểm tra index và tài liệu."
        )

    return RetrieverEvaluator.from_metric_names(
        METRICS,
        retriever=rag.router_retriever,
        node_postprocessors=rag.postprocessor_factory.build(),
    )


def extract_metric_values(result):
    """
    Lấy tên metric và điểm số từ kết quả evaluator.

    metric_vals_dict thường có dạng:
    {
        "hit_rate": ...,
        "mrr": ...,
        "precision": ...
    }
    """
    metric_values = {}

    for metric in METRICS:
        metric_result = result.metric_vals_dict.get(metric)

        if metric_result is None:
            metric_values[metric] = None
            continue

        # Một số phiên bản trả về EvaluationMetricResult,
        # một số trường hợp có thể trả về giá trị số.
        if hasattr(metric_result, "score"):
            score = metric_result.score
        else:
            score = metric_result

        metric_values[metric] = (
            float(score) if score is not None else None
        )

    return metric_values


def get_result_field(result, field, default=None):
    """Đọc thuộc tính kết quả một cách an toàn."""
    value = getattr(result, field, default)
    return default if value is None else value


def save_results(rows):
    """Xuất kết quả chi tiết và tổng hợp ra Excel."""

    RESULT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    # =========================
    # Sheet 1: Query Results
    # =========================
    detail_sheet = workbook.active
    detail_sheet.title = "Query Results"

    headers = [
        "query_id",
        "query",
        "expected_ids",
        "retrieved_ids",
        "hit_rate",
        "mrr",
        "precision",
        "status",
        "error",
    ]

    detail_sheet.append(headers)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
    )
    thin_gray = Side(
        style="thin",
        color="D9E2F3",
    )
    cell_border = Border(bottom=thin_gray)

    for cell in detail_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = cell_border

    metric_values = {
        metric: [] for metric in METRICS
    }

    for row in rows:
        detail_sheet.append([
            row["query_id"],
            row["query"],
            json.dumps(
                row["expected_ids"],
                ensure_ascii=False,
            ),
            json.dumps(
                row["retrieved_ids"],
                ensure_ascii=False,
            ),
            row["scores"]["hit_rate"],
            row["scores"]["mrr"],
            row["scores"]["precision"],
            row["status"],
            row["error"],
        ])

        for metric in METRICS:
            score = row["scores"].get(metric)

            if (
                row["status"] == "evaluated"
                and score is not None
            ):
                metric_values[metric].append(score)

    # =========================
    # Sheet 2: Summary
    # =========================
    summary_sheet = workbook.create_sheet("Summary")

    summary_sheet.append([
        "metric",
        "mean_score",
        "sample_count",
    ])

    for cell in summary_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = cell_border

    for metric, scores in metric_values.items():
        summary_sheet.append([
            metric,
            mean(scores) if scores else None,
            len(scores),
        ])

    # =========================
    # Định dạng Excel
    # =========================
    status_colors = {
        "evaluated": "E2F0D9",
        "error": "FCE4D6",
    }

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )
                cell.border = cell_border

        for column in sheet.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column
            )

            sheet.column_dimensions[
                column[0].column_letter
            ].width = min(
                max(max_length + 2, 14),
                55,
            )

        for row_idx in range(2, sheet.max_row + 1):
            for col_idx in range(1, sheet.max_column + 1):
                cell = sheet.cell(row_idx, col_idx)

                if (
                    sheet.title == "Query Results"
                    and cell.column == 8
                ):
                    color = status_colors.get(cell.value)

                    if color:
                        cell.fill = PatternFill(
                            fill_type="solid",
                            fgColor=color,
                        )

                if (
                    sheet.title == "Query Results"
                    and cell.column in (5, 6, 7)
                    and cell.value is not None
                ):
                    cell.number_format = "0.0000"

                if (
                    sheet.title == "Summary"
                    and cell.column == 2
                    and cell.value is not None
                ):
                    cell.number_format = "0.0000"

    detail_sheet.row_dimensions[1].height = 30
    summary_sheet.row_dimensions[1].height = 30

    workbook.save(RESULT_PATH)

    print(f"\nĐã lưu kết quả: {RESULT_PATH.resolve()}")
    print(f"Tổng số câu hỏi: {len(rows)}")

    evaluated_count = sum(
        row["status"] == "evaluated"
        for row in rows
    )
    error_count = sum(
        row["status"] == "error"
        for row in rows
    )

    print(f"Số câu đánh giá thành công: {evaluated_count}")
    print(f"Số câu bị lỗi: {error_count}")


async def main():
    if not DATASET_PATH.exists():
        raise FileNotFoundError(
            f"Không tìm thấy dataset: {DATASET_PATH}. "
            "Hãy chạy generate_dataset.py trước."
        )

    with DATASET_PATH.open(
        "r",
        encoding="utf-8",
    ) as file:
        dataset = json.load(file)

    queries = dataset.get("queries", {})
    relevant_docs = dataset.get("relevant_docs", {})

    if not queries:
        raise ValueError(
            "Dataset không có câu hỏi trong trường 'queries'."
        )

    rag = RAGService()
    evaluator = build_evaluator(rag)

    print(f"Số câu hỏi: {len(queries)}")
    print(f"Các metrics: {METRICS}")

    rows = []

    for index, (query_id, query) in enumerate(
        queries.items(),
        start=1,
    ):
        expected_ids = relevant_docs.get(query_id, [])

        row = {
            "query_id": query_id,
            "query": query,
            "expected_ids": expected_ids,
            "retrieved_ids": [],
            "scores": {
                metric: None for metric in METRICS
            },
            "status": "pending",
            "error": "",
        }

        print(f"\n[{index}/{len(queries)}] {query}")

        try:
            if not expected_ids:
                raise ValueError(
                    f"Query {query_id} không có expected_ids."
                )

            result = await evaluator.aevaluate(
                query=query,
                expected_ids=expected_ids,
            )

            # Đọc danh sách node thực tế được trả về.
            retrieved_ids = get_result_field(
                result,
                "retrieved_ids",
                [],
            )

            row["retrieved_ids"] = list(retrieved_ids)

            # Đọc điểm từng metric.
            row["scores"] = extract_metric_values(result)
            row["status"] = "evaluated"

            print(
                "  -> "
                + ", ".join(
                    f"{metric}={score:.4f}"
                    for metric, score in row["scores"].items()
                    if score is not None
                )
            )

        except Exception as exc:
            row["status"] = "error"
            row["error"] = str(exc)

            print(f"  -> Lỗi: {exc}")

        rows.append(row)

    save_results(rows)

    print("\nĐánh giá hoàn tất.")


if __name__ == "__main__":
    asyncio.run(main())
