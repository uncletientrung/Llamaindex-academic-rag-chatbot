
import json
import uuid
from pathlib import Path

from ragservice import RAGService
from evaluation.response_evaluation.evaluators import build_evaluators
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DATASET_PATH = Path("evaluation/response_evaluation/dataset.json")
RESULT_PATH = Path("evaluation/response_evaluation/results/evaluation.xlsx")


# Thứ tự các cột trong file Excel
COLUMNS = [
    # Thông tin định danh
    "eval_id",
    "query_id",
    "group",
    "query",
    "reference",

    # Kết quả RAG
    "response",
    "contexts",

    # Faithfulness
    "faithfulness_score",

    # Answer Relevancy
    "answer_relevancy_score",

    # Context Relevancy
    "context_relevancy_score",

    # Correctness
    "correctness_score",

    # Semantic Similarity
    "semantic_similarity_score",

    # Trạng thái
    "status",
    "error",
]


def serialize_value(value):
    """Chuyển dữ liệu về dạng phù hợp để ghi vào Excel."""
    if value is None:
        return ""

    if isinstance(value, (dict, list, tuple)):
        return json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
            default=str,
        )

    return value


def create_row(
    eval_id,
    query_id,
    query,
    reference,
    group,
):
    """Tạo dòng kết quả với đầy đủ các cột."""
    row = {column: "" for column in COLUMNS}

    row.update({
        "eval_id": eval_id,
        "query_id": query_id,
        "group": group,
        "query": query,
        "reference": reference or "",
        "status": "pending",
    })

    return row


def save_excel(rows):
    """Xuất Excel có định dạng và tự động điều chỉnh kích thước."""

    RESULT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Evaluation Results"

    # Màu sắc
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        name="Arial",
        bold=True,
        color="FFFFFF",
        size=11,
    )
    body_font = Font(
        name="Arial",
        size=10,
    )

    thin_gray = Side(
        style="thin",
        color="D9E2F3",
    )
    cell_border = Border(
        bottom=thin_gray,
    )

    # 1. Ghi tiêu đề
    for col_idx, column in enumerate(COLUMNS, start=1):
        cell = worksheet.cell(
            row=1,
            column=col_idx,
            value=column,
        )
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = cell_border

    # 2. Ghi dữ liệu
    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, column in enumerate(COLUMNS, start=1):
            value = serialize_value(row_data.get(column))

            cell = worksheet.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

            cell.font = body_font
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            cell.border = cell_border

            # Tô màu theo trạng thái
            if column == "status":
                if value == "evaluated":
                    cell.fill = PatternFill(
                        "solid",
                        fgColor="E2F0D9",
                    )
                elif value == "no_context":
                    cell.fill = PatternFill(
                        "solid",
                        fgColor="FFF2CC",
                    )
                elif value == "error":
                    cell.fill = PatternFill(
                        "solid",
                        fgColor="FCE4D6",
                    )

            # Định dạng các cột điểm số
            if column.endswith("_score") and isinstance(
                value, (int, float)
            ):
                cell.number_format = "0.00"

    # 3. Tự động điều chỉnh độ rộng cột
    for col_idx, column in enumerate(COLUMNS, start=1):
        col_letter = get_column_letter(col_idx)

        max_length = len(column)

        for row_idx in range(2, worksheet.max_row + 1):
            value = worksheet.cell(
                row=row_idx,
                column=col_idx,
            ).value

            if value is not None:
                lines = str(value).splitlines() or [""]
                max_length = max(
                    max_length,
                    max(len(line) for line in lines),
                )

        if column in (
            "query",
            "reference",
            "response",
            "contexts",
            "error",
        ):
            width = min(max(max_length * 0.9, 20), 55)
        else:
            width = min(max(max_length + 3, 12), 24)

        worksheet.column_dimensions[col_letter].width = width

    # 4. Tự động điều chỉnh chiều cao hàng
    for row_idx in range(2, worksheet.max_row + 1):
        max_lines = 1

        for col_idx in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(
                row=row_idx,
                column=col_idx,
            )

            if cell.value is None:
                continue

            col_width = worksheet.column_dimensions[
                get_column_letter(col_idx)
            ].width or 12

            text = str(cell.value)
            lines = text.splitlines() or [""]

            estimated_lines = sum(
                max(
                    1,
                    (
                        len(line) + int(col_width) - 1
                    ) // int(col_width),
                )
                for line in lines
            )

            max_lines = max(max_lines, estimated_lines)

        # Giới hạn chiều cao để tránh hàng quá lớn
        worksheet.row_dimensions[row_idx].height = min(
            max(20, max_lines * 15),
            180,
        )

    # 5. Cố định hàng tiêu đề và bật bộ lọc
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # 6. Định dạng chiều cao hàng tiêu đề
    worksheet.row_dimensions[1].height = 32

    # 7. Lưu file
    workbook.save(RESULT_PATH)

    print(f"\nĐã lưu Excel tại: {RESULT_PATH.resolve()}")
    print(f"Tổng số mẫu: {len(rows)}")


def main():
    rag = RAGService()
    evaluators = build_evaluators()

    with open(
        DATASET_PATH,
        "r",
        encoding="utf-8",
    ) as file:
        dataset = json.load(file)

    rows = []
    total = len(dataset)

    for i, item in enumerate(dataset, start=1):
        query_id = item["id"]
        eval_id = str(uuid.uuid4())
        query = item["query"]
        reference = item.get("reference")
        group = item.get("group", "all")

        row = create_row(
            eval_id=eval_id,
            query_id=query_id,
            query=query,
            reference=reference,
            group=group,
        )

        print(f"[{i}/{total}] {query}")

        try:
            # 1. Lấy câu trả lời từ RAG
            result = rag.answer_with_context(
                question=query,
                group=group,
            )

            contexts = result.get("contexts") or []
            response = result.get("response", "")

            row["response"] = response
            row["contexts"] = contexts

            # 2. Không có context
            if not contexts:
                row["status"] = "no_context"
                rows.append(row)

                print("  -> Không có context, bỏ qua đánh giá.")
                continue

            # 3. Faithfulness
            eval_result = evaluators["faithfulness"].evaluate(
                query=query,
                contexts=contexts,
                response=response,
            )
            row["faithfulness_score"] = eval_result.score

            # 4. Answer Relevancy
            eval_result = evaluators["answer_relevancy"].evaluate(
                query=query,
                contexts=contexts,
                response=response,
            )
            row["answer_relevancy_score"] = eval_result.score

            # 5. Context Relevancy
            eval_result = evaluators["context_relevancy"].evaluate(
                query=query,
                contexts=contexts,
            )
            row["context_relevancy_score"] = eval_result.score

            # 6. Các tiêu chí cần đáp án chuẩn
            if reference:
                # Correctness
                eval_result = evaluators["correctness"].evaluate(
                    query=query,
                    response=response,
                    reference=reference,
                )
                row["correctness_score"] = eval_result.score

                # Semantic Similarity
                eval_result = evaluators["semantic_similarity"].evaluate(
                    response=response,
                    reference=reference,
                )
                row["semantic_similarity_score"] = eval_result.score

            row["status"] = "evaluated"
            print("  -> Đánh giá hoàn tất.")

        except Exception as exc:
            row["status"] = "error"
            row["error"] = str(exc)

            print(f"  -> Lỗi: {exc}")

        # Luôn giữ lại dòng hiện tại, kể cả khi lỗi
        rows.append(row)

    # 7. Xuất kết quả
    save_excel(rows)


if __name__ == "__main__":
    main()
